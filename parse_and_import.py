
import json
import os
import re
import sqlite3

from docx import Document

# 读取文档
doc_path = os.path.join("C:\\Users\\Administrator\\Desktop", "试卷系统", "姐妹团考试原题.docx")
doc = Document(doc_path)

# 提取所有段落
texts = []
for para in doc.paragraphs:
    text = para.text.strip().replace('\xa0', ' ').replace('\u3000', ' ')
    if text:
        texts.append(text)

# 提取所有表格
tables = []
for table in doc.tables:
    rows = []
    for row in table.rows:
        cells = [cell.text.strip().replace('\xa0', ' ').replace('\u3000', ' ') for cell in row.cells]
        if cells and any(c for c in cells):
            rows.append(cells)
    if rows:
        tables.append(rows)

print(f"Total paragraphs: {len(texts)}")
print(f"Total tables: {len(tables)}")

# 改进的解析逻辑
questions = []
table_idx = 0

i = 0
while i < len(texts):
    text = texts[i]
    
    # 跳过标题和说明（前5行左右）
    if i < 5:
        i += 1
        continue
    
    # 检测题目编号行
    q_num_match = re.match(r'^(\d+)[\.\s]*$', text)
    q_inline_match = re.match(r'^(\d+)[\.\s]+(.+)', text)
    
    if q_num_match or q_inline_match:
        # 开始新题目
        if q_num_match:
            q_num = int(q_num_match.group(1))
            # 下一行应该是题目内容
            q_content = ""
            i += 1
            if i < len(texts):
                # 检查下一行是否是题型标记
                next_text = texts[i]
                if '[单选题]' in next_text or '[多选题]' in next_text or '[填空题]' in next_text or '【单选】' in next_text or '【多选】' in next_text:
                    # 题型标记行，题目内容就是这一行（去掉题型标记）
                    q_content = re.sub(r'\s*[\[【].*?[】\]].*', '', next_text).strip()
                    i += 1
                else:
                    q_content = next_text
                    i += 1
        else:
            q_num = int(q_inline_match.group(1))
            q_content = q_inline_match.group(2).strip()
            i += 1
        
        # 去除题型标记
        q_content = re.sub(r'\s*[\[【](单选题|多选题|填空题)[】\]].*', '', q_content).strip()
        q_content = re.sub(r'\s*\*$', '', q_content).strip()  # 去除末尾的 *
        
        q_type = 'unknown'
        options = []
        answer = ''
        blanks = []
        key_points = ''
        
        # 检查当前行或回溯的题型标记
        if i > 0 and i < len(texts):
            check_text = texts[i-1]
            if '[单选题]' in check_text or '【单选】' in check_text:
                q_type = 'single'
            elif '[多选题]' in check_text or '【多选】' in check_text:
                q_type = 'multiple'
            elif '[填空题]' in check_text:
                q_type = 'fill'
        
        # 如果题目行本身包含题型标记
        if q_num_match:
            if i > 1:
                check_text = texts[i-1]
                if '[单选题]' in check_text or '【单选】' in check_text:
                    q_type = 'single'
                elif '[多选题]' in check_text or '【多选】' in check_text:
                    q_type = 'multiple'
                elif '[填空题]' in check_text:
                    q_type = 'fill'
        else:
            # 内联题目，检查原始行
            raw = q_inline_match.group(0) if q_inline_match else ""
            if '[单选题]' in raw or '【单选】' in raw:
                q_type = 'single'
            elif '[多选题]' in raw or '【多选】' in raw:
                q_type = 'multiple'
            elif '[填空题]' in raw:
                q_type = 'fill'
        
        # 处理选择题选项
        if q_type in ('single', 'multiple') and table_idx < len(tables):
            opts = []
            correct_opts = []
            for row in tables[table_idx]:
                opt_text = row[0] if row else ''
                if '(正确答案)' in opt_text:
                    m = re.match(r'^([A-D])', opt_text)
                    if m:
                        correct_opts.append(m.group(1))
                # 清理选项文本
                clean_opt = opt_text.replace('(正确答案)', '')
                opts.append(clean_opt)
            options = opts
            answer = ','.join(correct_opts)
            table_idx += 1
        
        # 收集填空题答案和简答题解析
        if q_type == 'fill':
            # 收集后续行直到下一个题目开始
            parse_lines = []
            while i < len(texts):
                line = texts[i]
                # 检查是否是下一个题目
                if re.match(r'^(\d+)[\.\s]*', line) and not line.startswith('空') and not line.startswith('答案解析') and '[填空题]' not in line and '[单选题]' not in line and '[多选题]' not in line and '【单选】' not in line and '【多选】' not in line:
                    # 可能是下一个题目，但先确认不是下划线行
                    if not line.startswith('____') and not line.startswith('___') and not line.startswith('答案解析') and not line.startswith('空'):
                        break
                
                if line.startswith('空'):
                    # 解析空答案
                    b_matches = re.findall(r'空(\d+)答案:([^空]+)', line)
                    for b in b_matches:
                        blanks.append({'blank_no': int(b[0]), 'answer': b[1].strip()})
                    i += 1
                    continue
                
                if line.startswith('答案解析') or line.startswith('答案解析：'):
                    parse_text = line.replace('答案解析：', '').replace('答案解析:', '').strip()
                    if parse_text:
                        parse_lines.append(parse_text)
                    i += 1
                    # 继续收集解析内容
                    while i < len(texts):
                        next_line = texts[i]
                        if re.match(r'^(\d+)[\.\s]*', next_line) and not next_line.startswith('空') and not next_line.startswith('答案解析') and not next_line.startswith('____') and '[填空题]' not in next_line and '[单选题]' not in next_line and '[多选题]' not in next_line and '【单选】' not in next_line and '【多选】' not in next_line:
                            break
                        if next_line.startswith('空') or next_line.startswith('答案解析'):
                            break
                        parse_lines.append(next_line)
                        i += 1
                    continue
                
                if line.startswith('____') or line.startswith('___'):
                    i += 1
                    continue
                
                # 其他行也加入解析
                if line and not line.startswith('空'):
                    parse_lines.append(line)
                i += 1
            
            if parse_lines:
                key_points = '\n'.join(parse_lines)
            
            # 如果blanks不为空，是填空题；否则是简答题
            if blanks:
                q_type = 'fill'
            else:
                q_type = 'essay'
        
        questions.append({
            'id': q_num,
            'content': q_content,
            'type': q_type,
            'options': options,
            'answer': answer,
            'key_points': key_points,
            'blanks': blanks
        })
        continue
    
    i += 1

# 统计
stats = {
    'total': len(questions),
    'single': len([q for q in questions if q['type'] == 'single']),
    'multiple': len([q for q in questions if q['type'] == 'multiple']),
    'fill': len([q for q in questions if q['type'] == 'fill']),
    'essay': len([q for q in questions if q['type'] == 'essay']),
    'unknown': len([q for q in questions if q['type'] == 'unknown'])
}

print(f"\nStats: {stats}")

# 保存JSON
output_path = os.path.join("C:\\Users\\Administrator\\Desktop", "试卷系统", "questions_parsed.json")
with open(output_path, 'w', encoding='utf-8') as f:
    json.dump({'questions': questions, 'stats': stats}, f, ensure_ascii=False, indent=2)

# 创建SQLite数据库
db_path = os.path.join("C:\\Users\\Administrator\\Desktop", "试卷系统", "exam.db")
conn = sqlite3.connect(db_path)
cursor = conn.cursor()

# 创建表
cursor.execute('''
CREATE TABLE IF NOT EXISTS questions (
    id INTEGER PRIMARY KEY,
    q_number INTEGER NOT NULL,
    content TEXT NOT NULL,
    type TEXT NOT NULL,
    options TEXT,
    answer TEXT,
    key_points TEXT,
    blanks TEXT,
    difficulty INTEGER DEFAULT 1
)
''')

cursor.execute('DELETE FROM questions')

for q in questions:
    cursor.execute('''
    INSERT INTO questions (q_number, content, type, options, answer, key_points, blanks, difficulty)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    ''', (
        q['id'],
        q['content'],
        q['type'],
        json.dumps(q['options'], ensure_ascii=False) if q['options'] else None,
        q['answer'],
        q['key_points'] if q['key_points'] else None,
        json.dumps(q['blanks'], ensure_ascii=False) if q['blanks'] else None,
        1
    ))

conn.commit()
conn.close()

print(f"\nDatabase created at: {db_path}")
print(f"Total questions imported: {len(questions)}")

# 输出样本
for q in questions[:10]:
    print(f"\nQ{q['id']} [{q['type']}]: {q['content'][:60]}")
    if q['options']:
        print(f"  Options: {q['options']}")
    if q['answer']:
        print(f"  Answer: {q['answer']}")
    if q['blanks']:
        print(f"  Blanks: {q['blanks']}")
    if q['key_points']:
        print(f"  Parse: {q['key_points'][:60]}...")
