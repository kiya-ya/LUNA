# -*- coding: utf-8 -*-
"""解析导入的题目文件（docx / xlsx）"""
import sys
import json
import os
import re
import zipfile
import xml.etree.ElementTree as ET


def parse_docx(file_path):
    """用 XML 直接解析 Word 文档，提取题目（docx 库会丢失某些选项行）"""
    with zipfile.ZipFile(file_path, 'r') as z:
        xml_content = z.read('word/document.xml').decode('utf-8')
    
    tree = ET.fromstring(xml_content)
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    
    paragraphs = []
    for para in tree.findall('.//w:p', ns):
        texts = []
        for run in para.findall('.//w:t', ns):
            if run.text:
                texts.append(run.text)
        full_text = ''.join(texts).strip()
        if full_text:
            paragraphs.append(full_text)
    
    questions = []
    i = 0
    while i < len(paragraphs):
        text = paragraphs[i]
        match = re.match(r'^(\d+)[\.\s、．]+(.+)', text)
        if not match:
            i += 1
            continue
        
        content = match.group(2).strip()
        i += 1
        
        # 判断题型并清理标签
        q_type = 'essay'
        if '[单选题]' in content or '（单选）' in content or '（单选题）' in content:
            q_type = 'single'
            content = content.replace('[单选题]', '').replace('（单选）', '').replace('（单选题）', '').strip()
        elif '[多选题]' in content or '（多选）' in content or '（多选题）' in content:
            q_type = 'multiple'
            content = content.replace('[多选题]', '').replace('（多选）', '').replace('（多选题）', '').strip()
        elif '[填空题]' in content or '（填空）' in content or '（填空题）' in content:
            q_type = 'fill'
            content = content.replace('[填空题]', '').replace('（填空）', '').replace('（填空题）', '').strip()
        
        content = re.sub(r'\s*\*$', '', content).strip()
        
        # 收集后续行直到下一题或答案解析
        collected = []
        while i < len(paragraphs):
            line = paragraphs[i].strip()
            if re.match(r'^(\d+)[\.\s、．]+', line):
                break
            if line.startswith('___') or line.startswith('——') or line.startswith('____'):
                i += 1
                continue
            if line.startswith('答案解析：') or line == '答案解析：' or line.startswith('答案解析:'):
                break
            collected.append(line)
            i += 1
        
        # 解析选项和答案
        options = []
        answer = ''
        key_points = ''
        
        # 选项格式：A10元/h（无分隔符）、A、值班活跃（顿号）、A. 19:30（点+空格）、A:xxx（冒号）、A.陪档位活跃（点无空格）
        option_pattern = re.compile(r'^([A-D])\s*[:\s\.、．]?\s*(.+)')
        
        for line in collected:
            opt_match = option_pattern.match(line)
            if opt_match:
                letter = opt_match.group(1)
                opt_text = opt_match.group(2).strip()
                if '(正确答案)' in opt_text or '（正确答案）' in opt_text:
                    opt_text = opt_text.replace('(正确答案)', '').replace('（正确答案）', '').strip()
                    if answer:
                        answer += ',' + letter
                    else:
                        answer = letter
                options.append(opt_text)
        
        # 无字母前缀的选项（如第9题：450, 480, 680, 980）
        if not options and collected and q_type in ('single', 'multiple'):
            if 2 <= len(collected) <= 4:
                for idx, line in enumerate(collected):
                    letter = chr(65 + idx)
                    opt_text = line.strip()
                    if '(正确答案)' in opt_text or '（正确答案）' in opt_text:
                        opt_text = opt_text.replace('(正确答案)', '').replace('（正确答案）', '').strip()
                        if answer:
                            answer += ',' + letter
                        else:
                            answer = letter
                    options.append(opt_text)
        
        # 收集答案解析
        while i < len(paragraphs):
            line = paragraphs[i].strip()
            if re.match(r'^(\d+)[\.\s、．]+', line):
                break
            if line.startswith('答案解析：') or line == '答案解析：' or line.startswith('答案解析:'):
                if line == '答案解析：' or line == '答案解析:':
                    i += 1
                    kp_lines = []
                    while i < len(paragraphs):
                        next_line = paragraphs[i].strip()
                        if re.match(r'^(\d+)[\.\s、．]+', next_line):
                            break
                        if next_line.startswith('___') or next_line.startswith('——'):
                            i += 1
                            continue
                        kp_lines.append(next_line)
                        i += 1
                    key_points = '\n'.join(kp_lines)
                else:
                    key_points = re.sub(r'^答案解析[:：]\s*', '', line).strip()
                    i += 1
                break
            i += 1
        
        if options and q_type == 'essay':
            if answer and len(answer) > 1:
                q_type = 'multiple'
            else:
                q_type = 'single'
        
        questions.append({
            'content': content,
            'type': q_type,
            'options': options,
            'answer': answer,
            'key_points': key_points,
            'blanks': [],
            'enabled': True
        })
    
    return questions


def parse_xlsx(file_path):
    """解析 Excel 文件，提取题目"""
    from openpyxl import load_workbook
    wb = load_workbook(file_path)
    ws = wb.active
    
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    
    headers = [str(h).strip().lower() if h else '' for h in rows[0]]
    data_rows = rows[1:]
    
    col_map = {}
    for idx, h in enumerate(headers):
        if h in ('content', '题目', '题目内容', '题干', 'question', 'title', '题干内容'):
            col_map['content'] = idx
        elif h in ('type', '题型', '类型'):
            col_map['type'] = idx
        elif h in ('option_a', '选项a', '选项A', 'a'):
            col_map['option_a'] = idx
        elif h in ('option_b', '选项b', '选项B', 'b'):
            col_map['option_b'] = idx
        elif h in ('option_c', '选项c', '选项C', 'c'):
            col_map['option_c'] = idx
        elif h in ('option_d', '选项d', '选项D', 'd'):
            col_map['option_d'] = idx
        elif h in ('answer', '答案', '正确答案', 'correct'):
            col_map['answer'] = idx
        elif h in ('key_points', '解析', '答案解析', 'explanation'):
            col_map['key_points'] = idx
    
    if 'content' not in col_map and len(headers) > 0:
        col_map['content'] = 0
    if 'type' not in col_map and len(headers) > 1:
        col_map['type'] = 1
    if 'answer' not in col_map:
        for idx, h in enumerate(headers):
            if '答案' in h or 'answer' in h:
                col_map['answer'] = idx
                break
    
    questions = []
    for row in data_rows:
        if not row or not row[0]:
            continue
        content = str(row[col_map.get('content', 0)]).strip() if col_map.get('content', 0) < len(row) else ''
        if not content:
            continue
        
        type_val = str(row[col_map.get('type', 1)]).strip().lower() if col_map.get('type', 1) < len(row) else 'single'
        if type_val in ('单选', 'single', '单选题', 'radio', '1'):
            q_type = 'single'
        elif type_val in ('多选', 'multiple', '多选题', 'checkbox', '2'):
            q_type = 'multiple'
        elif type_val in ('填空', 'fill', '填空题', '3'):
            q_type = 'fill'
        elif type_val in ('简答', 'essay', '简答题', '问答题', '4'):
            q_type = 'essay'
        else:
            q_type = 'single'
        
        options = []
        for opt_key in ['option_a', 'option_b', 'option_c', 'option_d']:
            idx = col_map.get(opt_key, -1)
            if idx >= 0 and idx < len(row) and row[idx]:
                options.append(str(row[idx]).strip())
        
        answer = str(row[col_map.get('answer', 2)]).strip() if col_map.get('answer', 2) < len(row) else ''
        key_points = str(row[col_map.get('key_points', 3)]).strip() if col_map.get('key_points', 3) < len(row) else ''
        
        questions.append({
            'content': content,
            'type': q_type,
            'options': options,
            'answer': answer,
            'key_points': key_points,
            'blanks': [],
            'enabled': True
        })
    
    return questions


def main():
    if len(sys.argv) < 4:
        print('Usage: python parse_import.py <file_path> <file_type> <output_path>', file=sys.stderr)
        sys.exit(1)
    
    file_path = sys.argv[1]
    file_type = sys.argv[2]
    output_path = sys.argv[3]
    
    if not os.path.exists(file_path):
        print(json.dumps({'error': 'File not found'}), file=sys.stderr)
        sys.exit(1)
    
    try:
        if file_type == 'docx':
            questions = parse_docx(file_path)
        elif file_type == 'xlsx':
            questions = parse_xlsx(file_path)
        else:
            print(json.dumps({'error': 'Unsupported file type'}), file=sys.stderr)
            sys.exit(1)
        
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump({'questions': questions}, f, ensure_ascii=False, indent=2)
        
        print(f'Parsed {len(questions)} questions')
    except Exception as e:
        print(json.dumps({'error': str(e)}), file=sys.stderr)
        sys.exit(1)


if __name__ == '__main__':
    main()
